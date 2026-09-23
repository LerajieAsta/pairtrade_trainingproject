"""口試簡報產生器的共用版面元件（沿用 thesis/1150922.pptx 的視覺系統）。

座標一律以英吋為單位；投影片 13.333 × 7.5 in。
文字標記：**粗體**、^^強調色粗體^^、~~灰色~~。
"""
import copy
import re

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION,
                             XL_TICK_LABEL_POSITION)
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = "0B2A5B"
BLUE = "2E5A92"
ACCENT = "C4553B"
LIGHT = "E8EDF5"
GREY = "6B7280"
TEXT = "1A1A1A"
WHITE = "FFFFFF"
GREEN = "2E7D5B"
FONT = "Microsoft JhengHei"

SLIDE_W = 13.333
SLIDE_H = 7.5
LEFT = 0.70
CONTENT_W = 11.90


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


# ── 文字 ──────────────────────────────────────────────────────────────

_TOKEN = re.compile(r"(\*\*.+?\*\*|\^\^.+?\^\^|~~.+?~~)")


def _parse_runs(text):
    """把標記字串拆成 [(文字, 樣式)]，樣式 ∈ {None, 'b', 'accent', 'muted'}。"""
    out = []
    for part in _TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], "b"))
        elif part.startswith("^^"):
            out.append((part[2:-2], "accent"))
        elif part.startswith("~~"):
            out.append((part[2:-2], "muted"))
        else:
            out.append((part, None))
    return out


def _set_font(run, size, bold, color):
    f = run.font
    f.size = Pt(size)
    f.bold = bool(bold)
    f.color.rgb = rgb(color)
    f.name = FONT
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rpr, qn("a:ea"))
    ea.set("typeface", FONT)


def fill_text_frame(tf, paras, size=17, color=TEXT, bold=False, align="l",
                    line=1.25, space_after=6, accent=ACCENT, muted=GREY):
    """paras：字串或 dict(text, size, bold, color, align, after) 的串列。"""
    tf.word_wrap = True
    first = True
    for p_spec in paras:
        if isinstance(p_spec, str):
            p_spec = {"text": p_spec}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                       "r": PP_ALIGN.RIGHT}[p_spec.get("align", align)]
        p.line_spacing = p_spec.get("line", line)
        p.space_after = Pt(p_spec.get("after", space_after))
        psize = p_spec.get("size", size)
        pbold = p_spec.get("bold", bold)
        pcolor = p_spec.get("color", color)
        for txt, style in _parse_runs(p_spec["text"]):
            r = p.add_run()
            r.text = txt
            if style == "b":
                _set_font(r, psize, True, pcolor)
            elif style == "accent":
                _set_font(r, psize, True, accent)
            elif style == "muted":
                _set_font(r, psize, pbold, muted)
            else:
                _set_font(r, psize, pbold, pcolor)
    return tf


def text(slide, x, y, w, h, paras, size=17, color=TEXT, bold=False, align="l",
         line=1.25, space_after=6, anchor="t", accent=ACCENT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    fill_text_frame(tf, paras, size=size, color=color, bold=bold, align=align,
                    line=line, space_after=space_after, accent=accent)
    return tb


# ── 投影片骨架 ─────────────────────────────────────────────────────────

class Deck:
    def __init__(self, template_path):
        self.prs = Presentation(template_path)
        sld_ids = self.prs.slides._sldIdLst
        for sld_id in list(sld_ids):
            self.prs.part.drop_rel(sld_id.rId)
            sld_ids.remove(sld_id)
        # 原檔的講者備忘稿母片沒有任何版面配置區，新建的備忘稿頁會拿不到本文框；
        # 補上 python-pptx 預設母片的版面配置區。
        nm = self.prs.notes_master
        if not list(nm.placeholders):
            from pptx.oxml.slide import CT_NotesMaster
            dflt = CT_NotesMaster.new_default()
            tree = nm._element.cSld.spTree
            for sp in list(dflt.cSld.spTree):
                if sp.tag == qn("p:sp"):
                    tree.append(sp)
        self.layout = self.prs.slide_layouts[6]
        self.n = 0
        self.timeline = []   # (頁碼, 標題, 分鐘)

    def slide(self, dark=False):
        s = self.prs.slides.add_slide(self.layout)
        self.n += 1
        if dark:
            bg = s.background.fill
            bg.solid()
            bg.fore_color.rgb = rgb(NAVY)
        return s

    def save(self, path):
        self.prs.save(path)


def header(slide, label, dark=False):
    text(slide, LEFT, 0.42, CONTENT_W, 0.30, label, size=12, bold=True,
         color=(LIGHT if dark else GREY), space_after=2)


def title(slide, t, dark=False, size=30, y=0.72, h=0.90):
    text(slide, LEFT, y, CONTENT_W, h, t, size=size, bold=True,
         color=(WHITE if dark else NAVY), space_after=0)


def page_number(slide, n, dark=False):
    text(slide, 12.30, 6.95, 0.50, 0.30, str(n), size=10,
         color=(LIGHT if dark else GREY), align="r")


def refs(slide, items, y=6.62, w=11.30):
    if not items:
        return
    text(slide, LEFT, y, w, 0.3 * len(items), list(items), size=9, color=GREY,
         line=1.05, space_after=1)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip("\n")


# ── 方塊與圖形 ─────────────────────────────────────────────────────────

def _no_shadow(shape):
    sppr = shape._element.spPr
    if sppr.find(qn("a:effectLst")) is None:
        etree.SubElement(sppr, qn("a:effectLst"))


def box(slide, x, y, w, h, paras=None, fill=LIGHT, color=TEXT, size=14,
        bold=False, align="l", anchor="t", shape="round", line=None,
        inset=(0.22, 0.16), line_spacing=1.25, space_after=6, accent=ACCENT,
        radius=6000):
    kind = {"round": MSO_SHAPE.ROUNDED_RECTANGLE, "rect": MSO_SHAPE.RECTANGLE,
            "oval": MSO_SHAPE.OVAL, "chevron": MSO_SHAPE.CHEVRON,
            "pentagon": MSO_SHAPE.PENTAGON}[shape]
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == "round":
        shp.adjustments[0] = radius / 100000
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = rgb(line[0])
        shp.line.width = Pt(line[1])
    _no_shadow(shp)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Inches(inset[0])
    tf.margin_top = tf.margin_bottom = Inches(inset[1])
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    if paras:
        if isinstance(paras, str):
            paras = [paras]
        fill_text_frame(tf, paras, size=size, color=color, bold=bold,
                        align=align, line=line_spacing, space_after=space_after,
                        accent=accent)
    return shp


def callout(slide, x, y, w, h, heading, body, style="light", size=13.5,
            heading_size=15):
    """heading + 段落。style：light（淺藍底，標題強調色）／dark（深藍底）／plain（淺藍底，標題深藍）。"""
    if style == "dark":
        fill, hcol, bcol = NAVY, WHITE, LIGHT
    elif style == "plain":
        fill, hcol, bcol = LIGHT, NAVY, TEXT
    else:
        fill, hcol, bcol = LIGHT, ACCENT, TEXT
    paras = []
    if heading:
        paras.append({"text": heading, "size": heading_size, "bold": True,
                      "color": hcol, "after": 8})
    for b in (body if isinstance(body, list) else [body]):
        if isinstance(b, dict):
            b = dict(b)
            b.setdefault("color", bcol)
            b.setdefault("size", size)
            paras.append(b)
        else:
            paras.append({"text": b, "size": size, "color": bcol, "after": 6})
    return box(slide, x, y, w, h, paras, fill=fill, color=bcol,
               accent=(ACCENT if style != "dark" else "F2B8A8"))


def stat(slide, x, y, w, big, label, color=NAVY, big_size=34, label_size=13,
         dark=False):
    text(slide, x, y, w, 0.62, big, size=big_size, bold=True,
         color=(WHITE if dark else color), space_after=0)
    text(slide, x, y + 0.66, w, 0.5, label, size=label_size,
         color=(LIGHT if dark else GREY), space_after=0)


def line_seg(slide, x1, y1, x2, y2, color=GREY, width=1.25, dash=False,
             arrow=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1),
                                   Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def arrow(slide, x1, y1, x2, y2, color=GREY, width=1.75):
    return line_seg(slide, x1, y1, x2, y2, color=color, width=width, arrow=True)


def polyline(slide, pts, color=NAVY, width=2.25):
    """pts：[(x_in, y_in), ...]，畫一條不封閉的折線。"""
    fb = slide.shapes.build_freeform(Inches(pts[0][0]), Inches(pts[0][1]),
                                     scale=1.0)
    fb.add_line_segments([(Inches(px), Inches(py)) for px, py in pts[1:]],
                         close=False)
    shp = fb.convert_to_shape()
    shp.fill.background()
    shp.line.color.rgb = rgb(color)
    shp.line.width = Pt(width)
    _no_shadow(shp)
    return shp


def dot(slide, cx, cy, d=0.14, color=ACCENT):
    return box(slide, cx - d / 2, cy - d / 2, d, d, None, fill=color,
               shape="oval", inset=(0, 0))


# ── 表格 ─────────────────────────────────────────────────────────────

def table(slide, x, y, col_w, rows, num_cols=(), size=14, row_h=0.34,
          header_size=None, bold_cells=(), accent_cells=(), hl_rows=(),
          center_cols=(), wrap_h=None):
    """rows[0] 為表頭。num_cols 靠右；bold_cells/accent_cells 為 (r, c) 集合；
    hl_rows 以淡紅底標出。wrap_h：{row_index: 高度} 覆寫列高。"""
    nr, nc = len(rows), len(rows[0])
    total_w = sum(col_w)
    heights = [(wrap_h or {}).get(i, row_h) for i in range(nr)]
    gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(total_w),
                                Inches(sum(heights)))
    tbl = gf.table
    for c, w in enumerate(col_w):
        tbl.columns[c].width = Inches(w)
    for r in range(nr):
        tbl.rows[r].height = Inches(heights[r])
    hsize = header_size or size
    for r in range(nr):
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.margin_top = Emu(27432)
            cell.margin_bottom = Emu(27432)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            if r == 0:
                fill, col, b, sz = NAVY, WHITE, True, hsize
            else:
                fill = WHITE if r % 2 == 1 else LIGHT
                if r in hl_rows:
                    fill = "F6DED7"
                col, b, sz = TEXT, (r, c) in bold_cells, size
                if (r, c) in accent_cells:
                    col, b = ACCENT, True
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            align = "r" if (c in num_cols and r > 0) else (
                "c" if c in center_cols else "l")
            tf = cell.text_frame
            tf.word_wrap = True
            fill_text_frame(tf, [str(rows[r][c])], size=sz, color=col, bold=b,
                            align=align, line=1.1, space_after=0)
    return gf


# ── 圖表 ─────────────────────────────────────────────────────────────

def _chart_fonts(chart, size=12, color=TEXT):
    chart.font.size = Pt(size)
    chart.font.name = FONT
    chart.font.color.rgb = rgb(color)


def _kill_invert(chart):
    """PowerPoint 的「負值反轉」在 dPt 層各自預設為真，序列層關掉不夠，須逐點注入。"""
    cs = chart._chartSpace
    for ser in cs.iter(qn("c:ser")):
        inv = ser.find(qn("c:invertIfNegative"))
        if inv is None:
            inv = etree.SubElement(ser, qn("c:invertIfNegative"))
            # 須位於 spPr 之後、dPt 之前
            sppr = ser.find(qn("c:spPr"))
            if sppr is not None:
                sppr.addnext(inv)
        inv.set("val", "0")
        for dpt in ser.findall(qn("c:dPt")):
            if dpt.find(qn("c:invertIfNegative")) is None:
                el = etree.Element(qn("c:invertIfNegative"))
                el.set("val", "0")
                dpt.find(qn("c:idx")).addnext(el)


def bar_chart(slide, x, y, w, h, categories, values, colors, horizontal=False,
              number_format="0.00", label_size=12, axis_size=11,
              gridlines=True, val_min=None, val_max=None, gap=80, major_unit=None,
              axis_format=None):
    cd = CategoryChartData()
    cd.categories = categories
    cd.add_series("s", values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w),
                                Inches(h), cd)
    ch = gf.chart
    _chart_fonts(ch, axis_size)
    ch.has_title = False
    ch.has_legend = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.font.size = Pt(label_size)
    dl.font.bold = True
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    ser = plot.series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(c)
    va = ch.value_axis
    va.has_major_gridlines = gridlines
    if gridlines:
        va.major_gridlines.format.line.color.rgb = rgb("D9DEE7")
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(axis_size)
    if val_min is not None:
        va.minimum_scale = val_min
    if val_max is not None:
        va.maximum_scale = val_max
    if major_unit is not None:
        va.major_unit = major_unit
    if axis_format is not None:
        va.tick_labels.number_format = axis_format
        va.tick_labels.number_format_is_linked = False
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(axis_size)
    ca.format.line.color.rgb = rgb("9CA3AF")
    if min(values) < 0:
        ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    _kill_invert(ch)
    return gf


def line_chart(slide, x, y, w, h, categories, series, colors, y_min=None,
               y_max=None, axis_size=11, legend=True, number_format="#,##0"):
    cd = CategoryChartData()
    cd.categories = categories
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(x), Inches(y),
                                Inches(w), Inches(h), cd)
    ch = gf.chart
    _chart_fonts(ch, axis_size)
    ch.has_title = False
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(12)
    for i, s in enumerate(ch.plots[0].series):
        s.format.line.color.rgb = rgb(colors[i])
        s.format.line.width = Pt(2.5)
        s.smooth = False
        s.marker.style = None
        from pptx.enum.chart import XL_MARKER_STYLE
        s.marker.style = XL_MARKER_STYLE.NONE
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = rgb("D9DEE7")
    va.format.line.fill.background()
    va.tick_labels.number_format = number_format
    va.tick_labels.number_format_is_linked = False
    if y_min is not None:
        va.minimum_scale = y_min
    if y_max is not None:
        va.maximum_scale = y_max
    ch.category_axis.tick_labels.font.size = Pt(axis_size)
    ch.category_axis.format.line.color.rgb = rgb("9CA3AF")
    return gf
