"""把 thesis/NN_*.md 轉成符合元智大學研究所學位論文格式規範的 Word 檔。

用法（於 repo 根目錄）：
    Project/Scripts/python.exe tools/thesis_docx/build_docx.py [--out 輸出.docx] [--no-word]

流程：
  1. 前處理 Markdown：依規範排列（前置頁 → 本文 → 參考文獻 → 附錄）、加表號表名、公式編號。
  2. pandoc（quarto 內建）轉成 docx，樣式取自本程式產生的 reference.docx。
  3. python-docx：分節與頁碼（前置頁 i、ii…，本文 1、2…）、表格格線。
  4. Word（finalize.ps1）：表格依版面寬度調整、更新目錄與表目錄、另存 PDF。

格式依據（元智大學研究所學位論文格式規範條例）：
  A4；上 3.5 cm、左 4 cm、右 2 cm、下 2 cm；頁碼在版面底端 1 cm 處置中；
  摘要至表目錄以小寫羅馬數字編頁（書名頁、審定書計頁但不印），第一章起以阿拉伯數字編頁；
  章名置中、章名下留雙倍行距；表號及表名在表上方、全文連續編號並列表目錄；
  中文書名（論文名）粗體、英文書名與期刊名斜體。
字型比照同所前屆論文：中文標楷體、英文 Times New Roman、內文 12 點。
"""
import argparse
import copy
import glob
import os
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
THESIS = os.path.join(ROOT, "thesis")
sys.path.insert(0, HERE)
from captions import CAPTIONS  # noqa: E402

# ── 論文基本資料。標 TODO 的欄位我無法自現有檔案確認，Word 檔內以黃底標出。
META = {
    "title_zh": "結合配對來源組裝與動作空間設計之配對交易系統建構與實證：以 S&P 500 為例",
    "title_en": ("Construction and Empirical Investigation of a Pairs Trading System Combining "
                 "Pair-Source Assembly and Action-Space Design: Evidence from the S&P 500"),
    "student_zh": "李伯修",
    "student_en": "Bo-Siou Li",
    "advisor_zh": "李詩政 博士",
    "advisor_en": "Dr. Shih-Cheng Lee",
    "school_zh": "元智大學",
    "school_en": "Yuan Ze University",
    "dept_zh": "資訊管理在職碩士專班",                      # TODO：依簡報封面；請核對系所全銜
    "dept_en": "【系所英文全銜】",                          # TODO
    "degree_en": "【學位英文名稱，例：Master of Science】",    # TODO
    "date_zh": "中華民國 一一五 年 十 月",                  # TODO：提送年月
    "date_en": "October 2026",                              # TODO
    "place_en": "Chung-Li, Taiwan, Republic of China",
}
TODO_FIELDS = {"dept_zh", "dept_en", "degree_en", "date_zh", "date_en"}

FONT_ZH = "標楷體"
FONT_EN = "Times New Roman"
FONT_CODE = "Consolas"
SECT = "§§SECT§§"
PANDOC_FROM = "markdown+east_asian_line_breaks-smart-subscript-superscript-strikeout"

# 換頁一律由樣式的「段落前分頁」處理：獨立的分頁段落若落在滿頁之後，會多出一張空白頁。
TAB ='`<w:r><w:tab/></w:r>`{=openxml}'


def field(instr, placeholder):
    return ('```{=openxml}\n<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
            f'<w:r><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            f'<w:r><w:t>{placeholder}</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>\n```\n\n')


def div(style, *paras):
    return f'::: {{custom-style="{style}"}}\n' + "\n\n".join(paras) + "\n:::\n\n"


def m(key):
    """取基本資料；未確認的欄位以「待確認」字元樣式（黃底）標出。"""
    v = META[key]
    return f'[{v}]{{custom-style="待確認"}}' if key in TODO_FIELDS else v


# ════════════════════════════════════════════════════════════════════
# 1. Markdown 前處理
# ════════════════════════════════════════════════════════════════════

def read(name):
    with open(os.path.join(THESIS, name), encoding="utf-8") as f:
        return f.read()


def thesis_files():
    files = sorted(os.path.basename(p) for p in glob.glob(os.path.join(THESIS, "[0-9][0-9]_*.md")))
    abstract = [f for f in files if f.startswith("00_")]
    refs = [f for f in files if "參考文獻" in f]
    appendix = [f for f in files if "附錄" in f]
    body = [f for f in files if f not in abstract + refs + appendix]
    if len(abstract) != 1 or len(refs) != 1:
        raise SystemExit("thesis/ 下應各有一個摘要檔與參考文獻檔")
    return abstract[0], body, refs[0], appendix


class Counter:
    def __init__(self):
        self.table = 0
        self.eq = 0


TABLE_SEP = re.compile(r"^(>\s?)*\s*\|[\s:|-]+\|\s*$")

_WIDE = "　-〿＀-￯"        # 全形標點
_DASH = "—–…‥"
_HAN = "一-鿿"
_Q = r"(?:>[ ]?)*[ \t]*"                    # 引用區塊的行首標記、清單續行的縮排


def join_cjk(text):
    """原稿每行約 40 字硬換行。pandoc 的 east_asian_line_breaks 只在「兩側都是全形字」時
    不補空白；破折號與「英數＋全形標點」的換行仍會多出空白，這裡先把這幾種換行接起來。"""
    nxt = rf"(?=[{_HAN}{_WIDE}{_DASH}A-Za-z(]|\$(?!\$)|\*\*|`[^`])"
    text = re.sub(rf"(?<=[{_WIDE}{_DASH}])\n{_Q}{nxt}", "", text)
    text = re.sub(rf"(?<=[{_HAN}A-Za-z0-9)%\]$*`])\n{_Q}(?=[{_WIDE}{_DASH}])", "", text)
    return text


def _math_pipes(line):
    """表格列內，數學式中的 \\| 會被當成欄位分隔，改成 \\vert；並去掉結尾 $ 前的空白
    （pandoc 要求結尾的 $ 前不可有空白，否則不當成數學式）。"""
    def fix(mm):
        return "$" + mm.group(1).replace(r"\|", r"\vert{}").strip() + "$"
    return re.sub(r"\$([^$]*)\$", fix, line)


def prep(name, cnt):
    """單一章節檔的前處理：表名、公式編號、去掉分隔線。"""
    prefix = name[:2]
    caps = list(CAPTIONS.get(prefix, []))
    lines = join_cjk(read(name)).split("\n")
    out, i, fence, used = [], 0, False, 0
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            fence = not fence
            out.append(ln)
            i += 1
            continue
        if fence:
            out.append(ln)
            i += 1
            continue
        if ln.strip() == "---":
            i += 1
            continue
        bare = re.sub(r"^(>\s?)+", "", ln)
        quote = ln[:len(ln) - len(bare)]
        # 表格：加上表號與表名
        if bare.lstrip().startswith("|") and i + 1 < len(lines) and TABLE_SEP.match(lines[i + 1]) \
                and not (out and re.sub(r"^(>\s?)+", "", out[-1]).lstrip().startswith("|")):
            if used >= len(caps):
                raise SystemExit(f"{name}：表格比 captions.py 列的多（第 {used + 1} 張，約第 {i + 1} 行）")
            cnt.table += 1
            q = quote.rstrip() + " " if quote else ""
            out += [q.rstrip(), f'{q}::: {{custom-style="表標題"}}',
                    f"{q}表 {cnt.table}　{caps[used]}", f"{q}:::", q.rstrip()]
            used += 1
            while i < len(lines) and re.sub(r"^(>\s?)+", "", lines[i]).lstrip().startswith("|"):
                out.append(_math_pipes(lines[i]))
                i += 1
            out.append(q.rstrip())
            continue
        # 區塊公式：置中並於右側編號
        if ln.startswith("$$"):
            buf = [ln[2:]]
            while not buf[-1].rstrip().endswith("$$"):
                i += 1
                buf.append(lines[i])
            tex = " ".join(buf).rstrip()[:-2].strip()
            cnt.eq += 1
            out += ["", '::: {custom-style="公式"}', f"{TAB}${tex}${TAB}（{cnt.eq}）", ":::", ""]
            i += 1
            continue
        out.append(ln)
        i += 1
    if used != len(caps):
        raise SystemExit(f"{name}：表格 {used} 張，captions.py 列了 {len(caps)} 個表名")
    return "\n".join(out).strip() + "\n\n"


def abstract_parts(name):
    text = read(name)
    zh, en = text.split("# Abstract")
    zh = zh.replace("# 摘要", "").replace("\n---\n", "\n").strip()
    return zh, en.strip()


def references_md(name):
    text = join_cjk(read(name)).split("# 查證紀錄")[0]
    text = text[text.index("## "):]                      # 略過清單開頭的說明
    text = re.sub(r"\n---\s*$", "", text.strip())
    blocks = []
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if block.startswith("## "):
            blocks.append(block)
        elif block:
            blocks.append(div("參考文獻", block).strip())
    return "# 參考文獻\n\n" + "\n\n".join(blocks) + "\n\n"


def front_md(abstract_file):
    zh, en = abstract_parts(abstract_file)
    T = META
    s = ""
    # 封面
    s += div("封面校名", "元　智　大　學")
    s += div("封面系所", m("dept_zh"))
    s += div("封面學位", "碩　士　論　文")
    s += div("封面題目", T["title_zh"])
    s += div("封面英題", T["title_en"])
    s += div("封面作者首", f"研 究 生：{T['student_zh']}")
    s += div("封面作者", f"指導教授：{T['advisor_zh']}")
    s += div("封面日期", m("date_zh"))
    s += div("分節", SECT + "cover")
    # 書名頁
    s += div("書名頁題目", T["title_zh"])
    s += div("書名頁英題", T["title_en"])
    s += div("書名頁作者首", f"研 究 生：{T['student_zh']}　　　　　Student: {T['student_en']}")
    s += div("書名頁作者", f"指導教授：{T['advisor_zh']}　　　Advisor: {T['advisor_en']}")
    s += div("書名頁校名首", T["school_zh"])
    s += div("書名頁校名", m("dept_zh"))
    s += div("書名頁校名", "碩士論文")
    s += div("書名頁英文首", "A Thesis")
    s += div("書名頁英文", f"Submitted to {m('dept_en')}")
    s += div("書名頁英文", T["school_en"])
    s += div("書名頁英文", "in Partial Fulfillment of the Requirements")
    s += div("書名頁英文", "for the Degree of")
    s += div("書名頁英文", m("degree_en"))
    s += div("書名頁英文", m("date_en"))
    s += div("書名頁英文", T["place_en"])
    s += div("書名頁日期", m("date_zh"))
    # 審定書
    s += div("審定書標題", "論文口試委員會審定書")
    s += div("前置置中", '[（裝訂紙本時於此頁置入審定書影本；上傳電子檔時依規定保留空白頁）]{custom-style="待確認"}')
    s += div("分節", SECT + "title")
    # 中文摘要
    s += div("摘要題目", T["title_zh"])
    s += div("前置置中", f"學生：{T['student_zh']}　　　　　　指導教授：{T['advisor_zh']}")
    s += div("前置置中", f"{T['school_zh']}　{m('dept_zh')}")
    s += div("前置標題", "摘要") + div("摘要內文", join_cjk(zh))
    # 英文摘要
    s += div("摘要題目", T["title_en"])
    s += div("前置置中", f"Student: {T['student_en']}　　　　Advisor: {T['advisor_en']}")
    s += div("前置置中", m("dept_en"))
    s += div("前置置中", T["school_en"])
    s += div("前置標題", "Abstract") + div("摘要內文", en)
    # 誌謝
    s += "# 誌謝\n\n" + '[（請於此撰寫誌謝辭）]{custom-style="待確認"}\n\n'
    # 目錄、表目錄
    s += "# 目錄\n\n" + field(r'TOC \o "1-2" \h \z \u', "（請在 Word 內按 F9 更新目錄）")
    s += "# 表目錄\n\n" + field(r'TOC \h \z \t "表標題,1"', "（請在 Word 內按 F9 更新表目錄）")
    s += div("分節", SECT + "front")
    return s


def build_markdown():
    abstract, body, refs, appendix = thesis_files()
    cnt = Counter()
    parts = [front_md(abstract)]
    chapters = [prep(f, cnt) for f in body] + [references_md(refs)] + [prep(f, cnt) for f in appendix]
    parts.append("".join(chapters))
    return "".join(parts), cnt


# ════════════════════════════════════════════════════════════════════
# 2. reference.docx：樣式
# ════════════════════════════════════════════════════════════════════

def _fonts(rpr, zh=FONT_ZH, en=FONT_EN):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:cs"), en)
    rf.set(qn("w:eastAsia"), zh)


def _style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH, base="Normal"):
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, kind)
        if base:
            st.base_style = doc.styles[base]
    return st


def _para(st, size=12, bold=None, align=None, before=0, after=0, first=None, left=None,
          line=20, keep=False, zh=FONT_ZH, en=FONT_EN, new_page=False, outline=None):
    f = st.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = False
    f.color.rgb = RGBColor(0, 0, 0)
    _fonts(st.element.get_or_add_rPr(), zh, en)
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(first) if first is not None else Pt(0)
    pf.left_indent = Pt(left) if left is not None else Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    pf.line_spacing = Pt(line)
    pf.keep_with_next = keep
    pf.page_break_before = new_page
    if outline is not None:                      # 讓非標題樣式也列入目錄
        ppr = st.element.get_or_add_pPr()
        ol = ppr.find(qn("w:outlineLvl"))
        if ol is None:
            ol = OxmlElement("w:outlineLvl")
            ppr.append(ol)
        ol.set(qn("w:val"), str(outline))
    return st


def make_reference(path, quarto):
    with open(path, "wb") as f:
        subprocess.run([quarto, "pandoc", "--print-default-data-file", "reference.docx"],
                       stdout=f, check=True)
    doc = Document(path)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(3.5), Cm(2.0)
    sec.left_margin, sec.right_margin = Cm(4.0), Cm(2.0)
    sec.header_distance, sec.footer_distance = Cm(1.5), Cm(1.0)
    sec.gutter = Cm(0)

    # 文件預設：字型與語言（zh-TW 才會套用避頭尾）
    rpr = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _fonts(rpr)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-TW")
    # 行尾標點不得懸掛到右邊界外（Word 預設允許，會吃掉規範要求的 2 cm 右邊界）
    dd = doc.styles.element.find(qn("w:docDefaults"))
    ppd = dd.find(qn("w:pPrDefault"))
    if ppd is None:
        ppd = OxmlElement("w:pPrDefault")
        dd.append(ppd)
    ppr0 = ppd.find(qn("w:pPr"))
    if ppr0 is None:
        ppr0 = OxmlElement("w:pPr")
        ppd.append(ppr0)
    op = OxmlElement("w:overflowPunct")
    op.set(qn("w:val"), "0")
    ppr0.insert(0, op)

    C, J, L = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT
    _para(doc.styles["Normal"], align=L)
    _para(_style(doc, "Body Text"), align=J, first=24, after=6)
    _para(_style(doc, "First Paragraph", base="Body Text"), align=J, first=24, after=6)
    _para(_style(doc, "Compact", base="Body Text"), align=J, after=3)
    _para(doc.styles["Heading 1"], size=18, bold=True, align=C, after=24, line=28, keep=True,
          new_page=True)
    _para(doc.styles["Heading 2"], size=16, bold=True, align=L, before=18, after=8, line=24, keep=True)
    _para(doc.styles["Heading 3"], size=14, bold=True, align=L, before=12, after=6, line=22, keep=True)
    bt = _para(_style(doc, "Block Text"), align=J, left=24, after=6)
    ppr = bt.element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    for k, v in (("val", "single"), ("sz", "12"), ("space", "8"), ("color", "808080")):
        left.set(qn("w:" + k), v)
    bdr.append(left)
    ppr.insert(0, bdr)
    _para(_style(doc, "Footnote Text"), size=10, line=14)
    _para(_style(doc, "Source Code"), size=10.5, zh=FONT_ZH, en=FONT_CODE, after=6, left=24)
    vc = _style(doc, "Verbatim Char", WD_STYLE_TYPE.CHARACTER, base=None)
    vc.font.size = Pt(10.5)
    _fonts(vc.element.get_or_add_rPr(), FONT_ZH, FONT_CODE)

    # 自訂樣式
    _para(_style(doc, "表標題"), align=C, before=10, after=4, keep=True)
    _para(_style(doc, "表格內文"), size=10.5, align=L, line=15)
    eq = _para(_style(doc, "公式"), align=L, before=6, after=6)
    eq.paragraph_format.tab_stops.add_tab_stop(Cm(7.5), WD_TAB_ALIGNMENT.CENTER)
    eq.paragraph_format.tab_stops.add_tab_stop(Cm(15.0), WD_TAB_ALIGNMENT.RIGHT)
    _para(_style(doc, "參考文獻"), align=L, left=24, first=-24, after=6)   # 靠左：含網址的行左右對齊會被撐開
    _para(_style(doc, "分節"), size=2, line=2)
    _para(_style(doc, "前置置中"), align=C, after=4)
    _para(_style(doc, "摘要題目"), size=14, bold=True, align=C, after=10, line=22, new_page=True)
    _para(_style(doc, "前置標題"), size=18, bold=True, align=C, before=4, after=18, line=28,
          keep=True, outline=0)
    _para(_style(doc, "審定書標題"), size=18, bold=True, align=C, before=6, after=10, line=28,
          new_page=True)
    _para(_style(doc, "封面校名"), size=28, bold=True, align=C, before=10, after=14, line=36)
    _para(_style(doc, "封面系所"), size=20, align=C, after=14, line=28)
    _para(_style(doc, "封面學位"), size=22, bold=True, align=C, after=0, line=30)
    _para(_style(doc, "封面題目"), size=20, bold=True, align=C, before=70, after=14, line=30)
    _para(_style(doc, "封面英題"), size=16, bold=True, align=C, after=0, line=24)
    _para(_style(doc, "封面作者首"), size=18, align=C, before=80, after=6, line=28)
    _para(_style(doc, "封面作者"), size=18, align=C, after=0, line=28)
    _para(_style(doc, "封面日期"), size=18, align=C, before=80, line=28)
    _para(_style(doc, "書名頁題目"), size=18, bold=True, align=C, before=6, after=10, line=28)
    _para(_style(doc, "書名頁英題"), size=14, bold=True, align=C, after=0, line=22)
    _para(_style(doc, "書名頁作者首"), size=14, align=C, before=30, after=4, line=22)
    _para(_style(doc, "書名頁作者"), size=14, align=C, after=0, line=22)
    _para(_style(doc, "書名頁校名首"), size=16, align=C, before=30, line=26)
    _para(_style(doc, "書名頁校名"), size=16, align=C, line=26)
    _para(_style(doc, "書名頁英文首"), size=12, align=C, before=30, line=20)
    _para(_style(doc, "書名頁英文"), size=12, align=C, line=20)
    _para(_style(doc, "書名頁日期"), size=14, align=C, before=30, line=22)
    _para(_style(doc, "摘要內文"), align=J, first=24, after=3, line=16.5)   # 摘要以一頁為限
    # 「待確認」：黃底。Word 的樣式不支援螢光筆（highlight），用網底（shd）
    todo = _style(doc, "待確認", WD_STYLE_TYPE.CHARACTER, base=None)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "FFFF00")
    todo.element.get_or_add_rPr().append(shd)
    doc.save(path)


# ════════════════════════════════════════════════════════════════════
# 3. 後處理：分節、頁碼、表格
# ════════════════════════════════════════════════════════════════════

def _sect_child(sect, tag, **attrs):
    """依 CT_SectPr 的子元素順序放入 type／pgNumType。"""
    order = ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz",
             "pgMar", "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt",
             "vAlign", "noEndnote", "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid"]
    el = sect.find(qn("w:" + tag))
    if el is None:
        el = OxmlElement("w:" + tag)
        idx = order.index(tag)
        after = [c for c in sect if c.tag.split("}")[1] in order[:idx]]
        if after:
            after[-1].addnext(el)
        else:
            sect.insert(0, el)
    for a in list(el.attrib):
        del el.attrib[a]
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    return el


def _page_field(par):
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(12)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text.strip() == "PAGE":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        else:
            r.text = text


def _borders(parent, tag, spec):
    el = parent.find(qn("w:" + tag))
    if el is not None:
        parent.remove(el)
    el = OxmlElement("w:" + tag)
    for side, sz in spec:
        b = OxmlElement("w:" + side)
        b.set(qn("w:val"), "single" if sz else "nil")
        if sz:
            b.set(qn("w:sz"), str(sz))
            b.set(qn("w:space"), "0")
            b.set(qn("w:color"), "000000")
        el.append(b)
    return el


TEXT_W = 8504          # 版面寬 15 cm（twips）


def _vis(s):
    return sum(2 if ord(c) > 0x2E7F else 1 for c in s)


def _unbreakable(txt):
    """儲存格內最長的不可斷字串（半形字數）。中文逐字可斷；英數以空白為界；
    不含中文的短內容（數字、區間、代號）整格視為不可斷。"""
    if _vis(txt) <= 18 and not re.search(r"[一-鿿]", txt):
        return _vis(txt)
    toks = re.split(r"[\s⺀-￿]+", txt)
    return max([len(x) for x in toks] + [2])


def _table_widths(t):
    """依內容分配欄寬。放得下就用自然寬度（表格比版面窄、置中）；放不下時，
    先給每欄「放得下最長不可斷字串」的寬度，剩餘寬度再按各欄還缺多少分配。"""
    rows = t._tbl.findall(qn("w:tr"))
    ncol = max(len(r.findall(qn("w:tc"))) for r in rows)
    vis, nice, strict, area = [2] * ncol, [6] * ncol, [4] * ncol, [1] * ncol
    for ri, r in enumerate(rows):
        for c, tc in enumerate(r.findall(qn("w:tc"))):
            txt = "".join(x.text or "" for x in tc.iter() if x.tag in (qn("w:t"), qn("m:t")))
            v = _vis(txt)
            vis[c] = max(vis[c], v)
            area[c] += v
            nice[c] = max(nice[c], _unbreakable(txt), min(v, 12) if ri == 0 else 0)
            strict[c] = max(strict[c], max([len(x) for x in re.split(r"[\s⺀-￿]+", txt)] + [0]))
    unit, pad = 112, 260                        # 10.5 點：半形字約 105 twips；儲存格左右邊距
    nat = [v * unit + pad for v in vis]
    if sum(nat) <= TEXT_W:
        ws = nat
    else:
        lo = [min(a * unit + pad, n) for a, n in zip(nice, nat)]
        if sum(lo) > TEXT_W * 0.7:              # 太擠：表頭與短數字也允許換行
            lo = [min(a * unit + pad, n) for a, n in zip(strict, nat)]
        if sum(lo) >= TEXT_W:
            ws = [int(TEXT_W * x / sum(lo)) for x in lo]
        else:                                   # 剩餘寬度按各欄文字量分配，不超過自然寬度
            ws, todo = list(lo), set(range(ncol))
            extra = TEXT_W - sum(lo)
            while extra > 1 and todo:
                tot = sum(area[i] for i in todo)
                full = [i for i in todo if ws[i] + extra * area[i] / tot >= nat[i]]
                if not full:
                    for i in todo:
                        ws[i] += int(extra * area[i] / tot)
                    break
                for i in full:
                    extra -= nat[i] - ws[i]
                    ws[i] = nat[i]
                    todo.remove(i)
    tblPr = t._tbl.tblPr
    for tag, attrs in (("tblW", {"w": sum(ws), "type": "dxa"}), ("tblLayout", {"type": "fixed"})):
        el = tblPr.find(qn("w:" + tag))
        if el is None:
            el = OxmlElement("w:" + tag)
            tblPr.append(el)
        for a in list(el.attrib):
            del el.attrib[a]
        for k, v in attrs.items():
            el.set(qn("w:" + k), str(v))
    grid = t._tbl.find(qn("w:tblGrid"))
    for gc in list(grid):
        grid.remove(gc)
    for w in ws:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for r in rows:
        for c, tc in enumerate(r.findall(qn("w:tc"))):
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.insert(0, tcW)
            tcW.set(qn("w:w"), str(ws[c]))
            tcW.set(qn("w:type"), "dxa")


def postprocess(src, dst):
    doc = Document(src)
    body = doc.element.body
    final = body.find(qn("w:sectPr"))
    for ref in final.findall(qn("w:headerReference")) + final.findall(qn("w:footerReference")):
        final.remove(ref)

    # ── 分節
    kinds = []
    for p in list(doc.paragraphs):
        if p.text.startswith(SECT):
            kinds.append(p.text[len(SECT):])
            for r in list(p._p.findall(qn("w:r"))):
                p._p.remove(r)
            sp = copy.deepcopy(final)
            p._p.get_or_add_pPr().append(sp)
    if kinds != ["cover", "title", "front"]:
        raise SystemExit(f"分節標記不符：{kinds}")
    sects = [s._sectPr for s in doc.sections]
    for sp in sects:
        _sect_child(sp, "type", val="nextPage")
    _sect_child(sects[0], "pgNumType", fmt="lowerRoman", start=1)   # 封面：不印頁碼
    _sect_child(sects[1], "pgNumType", fmt="lowerRoman", start=1)   # 書名頁、審定書：計頁不印
    _sect_child(sects[2], "pgNumType", fmt="lowerRoman")            # 摘要至表目錄：接續
    _sect_child(sects[3], "pgNumType", fmt="decimal", start=1)      # 本文、參考文獻、附錄
    for i, sec in enumerate(doc.sections):
        sec.header.is_linked_to_previous = False
        ft = sec.footer
        ft.is_linked_to_previous = False
        par = ft.paragraphs[0]
        for r in list(par.runs):
            r._r.getparent().remove(r._r)
        if i >= 2:
            _page_field(par)

    # ── 表格：上下粗線、表頭下細線、列間細線；表頭跨頁重複；內文 10.5 點
    for t in doc.tables:
        _table_widths(t)
        tblPr = t._tbl.tblPr
        tblPr.append(_borders(tblPr, "tblBorders", [("top", 12), ("left", 0), ("bottom", 12),
                                                     ("right", 0), ("insideH", 4), ("insideV", 0)]))
        jc = tblPr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            tblPr.append(jc)
        jc.set(qn("w:val"), "center")
        for ri, row in enumerate(t.rows):
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:        # 同一列不跨頁
                trPr.insert(0, OxmlElement("w:cantSplit"))
            if ri == 0 and trPr.find(qn("w:tblHeader")) is None:
                trPr.append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                if ri == 0:
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcPr.append(_borders(tcPr, "tcBorders", [("bottom", 8)]))
                for par in cell.paragraphs:
                    par.style = doc.styles["表格內文"]
                    if ri == 0:
                        for r in par.runs:
                            r.font.bold = True
    doc.save(dst)
    return len(doc.tables)


# ════════════════════════════════════════════════════════════════════

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(THESIS, "論文全文.docx"))
    ap.add_argument("--no-word", action="store_true", help="略過 Word 的更新目錄與 PDF 輸出")
    ap.add_argument("--keep", action="store_true", help="保留中間檔（印出暫存目錄）")
    a = ap.parse_args()
    quarto = shutil.which("quarto")
    if not quarto:
        raise SystemExit("找不到 quarto（需要它內建的 pandoc）")

    tmp = tempfile.mkdtemp(prefix="thesis_docx_")
    md, cnt = build_markdown()
    md_path = os.path.join(tmp, "thesis.md")
    with open(md_path, "w", encoding="utf-8", newline="\n") as f:
        f.write(md)
    ref = os.path.join(tmp, "reference.docx")
    make_reference(ref, quarto)
    raw = os.path.join(tmp, "raw.docx")
    r = subprocess.run([quarto, "pandoc", md_path, "-f", PANDOC_FROM, "-t", "docx",
                        "--reference-doc", ref, "-o", raw], capture_output=True, text=True,
                       encoding="utf-8")
    if r.stderr.strip():
        print("pandoc 訊息：\n" + r.stderr.strip())
    if r.returncode:
        raise SystemExit("pandoc 失敗")
    final = os.path.join(tmp, "final.docx")
    ntab = postprocess(raw, final)
    if ntab != cnt.table:
        print(f"注意：Word 內表格 {ntab} 張，表名 {cnt.table} 個")
    if not a.no_word:
        ps = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                             os.path.join(HERE, "finalize.ps1"), "-Path", final],
                            capture_output=True, text=True)
        print(ps.stdout.strip())
        if ps.returncode:
            print(ps.stderr.strip())
            raise SystemExit("Word 後處理失敗")
    out = os.path.abspath(a.out)
    shutil.copyfile(final, out)
    pdf = final[:-5] + ".pdf"
    if os.path.exists(pdf):
        shutil.copyfile(pdf, out[:-5] + ".pdf")
    print(f"表 {cnt.table} 張、編號公式 {cnt.eq} 式 → {out}")
    if a.keep:
        print("中間檔：", tmp)
    else:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
